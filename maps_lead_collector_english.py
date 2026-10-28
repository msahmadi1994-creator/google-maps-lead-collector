import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import threading
import queue
import sqlite3
import csv
import json
import re
import time
from pathlib import Path
from urllib.parse import urlparse, urlunparse, quote_plus


APP_DIR = Path(__file__).resolve().parent
DB_FILE = APP_DIR / "maps_leads.sqlite3"

FIELDS = {
    "name": "Business Name",
    "phone": "Phone",
    "address": "Address",
    "website": "Website",
    "email": "Email",
    "category": "Category",
    "rating": "Rating",
    "reviews": "Review Count",
    "hours": "Opening Hours",
    "maps_url": "Google Maps URL",
    "latitude": "Latitude",
    "longitude": "Longitude",
}

def clean(value):
    return re.sub(r"\s+", " ", value or "").strip()

def normalize_url(url):
    """Canonicalize Google Maps URLs for deduplication."""
    if not url:
        return ""
    url = url.strip()
    if url.startswith("/"):
        url = "https://www.google.com" + url
    p = urlparse(url)
    if not p.scheme or not p.netloc:
        return ""
    return urlunparse((p.scheme.lower(), p.netloc.lower(), p.path.rstrip("/"), "", "", ""))

def init_db():
    with sqlite3.connect(DB_FILE) as con:
        con.execute("""
            CREATE TABLE IF NOT EXISTS leads (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                maps_url TEXT UNIQUE,
                name TEXT DEFAULT '',
                phone TEXT DEFAULT '',
                address TEXT DEFAULT '',
                website TEXT DEFAULT '',
                email TEXT DEFAULT '',
                category TEXT DEFAULT '',
                rating TEXT DEFAULT '',
                reviews TEXT DEFAULT '',
                hours TEXT DEFAULT '',
                latitude TEXT DEFAULT '',
                longitude TEXT DEFAULT '',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)
        con.commit()

def save_record(record):
    """Insert new records; fill previously empty fields on existing records."""
    url = normalize_url(record.get("maps_url", ""))
    if not url:
        return "skipped"

    record["maps_url"] = url
    columns = [
        "maps_url", "name", "phone", "address", "website", "email",
        "category", "rating", "reviews", "hours", "latitude", "longitude"
    ]

    with sqlite3.connect(DB_FILE, timeout=30) as con:
        old = con.execute(
            "SELECT " + ", ".join(columns) + " FROM leads WHERE maps_url=?",
            (url,)
        ).fetchone()

        if old:
            old_data = dict(zip(columns, old))
            updates = {}
            for key in columns:
                new_value = clean(str(record.get(key, "") or ""))
                old_value = clean(str(old_data.get(key, "") or ""))
                # Keep existing data unless the field was empty.
                if key != "maps_url" and not old_value and new_value:
                    updates[key] = new_value

            if updates:
                assignments = ", ".join(f"{k}=?" for k in updates)
                con.execute(
                    f"UPDATE leads SET {assignments}, updated_at=CURRENT_TIMESTAMP WHERE maps_url=?",
                    list(updates.values()) + [url]
                )
                con.commit()
                return "updated"
            return "duplicate"

        values = [record.get(k, "") for k in columns]
        placeholders = ",".join("?" for _ in columns)
        con.execute(
            f"INSERT INTO leads ({','.join(columns)}) VALUES ({placeholders})",
            values
        )
        con.commit()
        return "new"

def get_text(page, selectors):
    for selector in selectors:
        try:
            loc = page.locator(selector).first
            if loc.count():
                value = clean(loc.inner_text(timeout=1500))
                if not value:
                    value = clean(loc.get_attribute("aria-label", timeout=1000) or "")
                if value:
                    return value
        except Exception:
            pass
    return ""

def get_attr(page, selector, attr="href"):
    try:
        loc = page.locator(selector).first
        if loc.count():
            return clean(loc.get_attribute(attr, timeout=1500) or "")
    except Exception:
        pass
    return ""

def extract_emails(text):
    found = re.findall(
        r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
        text or ""
    )
    return ", ".join(sorted(set(found)))

def extract_details(page, fallback_name=""):
    # Wait for a place details panel, but do not hang indefinitely.
    try:
        page.locator("h1").first.wait_for(timeout=8000)
    except Exception:
        pass

    record = {
        "name": get_text(page, ["h1.DUwDvf", "h1"]),
        "phone": "",
        "address": "",
        "website": "",
        "email": "",
        "category": "",
        "rating": "",
        "reviews": "",
        "hours": "",
        "maps_url": normalize_url(page.url),
        "latitude": "",
        "longitude": "",
    }

    if not record["name"]:
        record["name"] = fallback_name

    # Address and phone are most reliably exposed through data-item-id.
    record["address"] = get_text(page, [
        'button[data-item-id="address"]',
        '[data-item-id="address"]',
        'button[aria-label*="Address"]',
        'button[aria-label*="Address"]',
        'button[aria-label*="Address"]',
    ])

    phone = get_attr(page, 'button[data-item-id^="phone:"]', "aria-label")
    if not phone:
        phone = get_text(page, [
            'button[data-item-id^="phone:"]',
            '[data-item-id^="phone:"]',
            'button[aria-label*="Phone"]',
            'button[aria-label*="Phone"]',
        ])

    # Remove labels such as "Phone: " from the extracted number.
    phone = re.sub(
        r"^(Phone|โทรศัพท์|电话|電話|โทร|تلفن|شماره تلفن)\s*:\s*",
        "", phone, flags=re.I
    )
    record["phone"] = clean(phone)

    # Website link: use Maps' website data item first.
    record["website"] = (
        get_attr(page, 'a[data-item-id="authority"]', "href")
        or get_attr(page, 'a[aria-label*="Website"]', "href")
        or get_attr(page, 'a[aria-label*="website"]', "href")
        or get_attr(page, 'a[aria-label*="网站"]', "href")
    )

    record["category"] = get_text(page, [
        'button[jsaction*="category"]',
        'button.DkEaL',
    ])

    record["rating"] = get_attr(page, 'div.F7nice span[aria-hidden="true"]', "aria-label")
    if not record["rating"]:
        record["rating"] = get_text(page, ['div.F7nice span[aria-hidden="true"]'])

    reviews = get_attr(page, 'div.F7nice span[aria-label*="reviews"]', "aria-label")
    if not reviews:
        reviews = get_attr(page, 'div.F7nice span[aria-label*="review"]', "aria-label")
    record["reviews"] = reviews or ""

    record["hours"] = get_text(page, [
        'div[aria-label*="Hours"]',
        'button[aria-label*="Hours"]',
        'table.eK4R0e',
    ])

    # Try coordinates from the Maps URL.
    match = re.search(r"!3d(-?\d+(?:\.\d+)?)!4d(-?\d+(?:\.\d+)?)", page.url)
    if match:
        record["latitude"] = match.group(1)
        record["longitude"] = match.group(2)
    else:
        match = re.search(r"/@(-?\d+(?:\.\d+)?),(-?\d+(?:\.\d+)?),", page.url)
        if match:
            record["latitude"] = match.group(1)
            record["longitude"] = match.group(2)

    # Email isn't normally shown by Maps; inspect visible page text as a fallback.
    try:
        record["email"] = extract_emails(page.locator("body").inner_text(timeout=2000))
    except Exception:
        pass

    return record

class MapsApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Google Maps Lead Collector")
        self.root.geometry("1200x720")
        self.events = queue.Queue()
        self.running = False
        self.stop_event = threading.Event()
        self.selected_fields = {}
        self._build_ui()
        init_db()
        self.refresh_table()
        self.root.after(150, self.process_events)

    def _build_ui(self):
        top = ttk.Frame(self.root, padding=8)
        top.pack(fill="x")

        ttk.Label(top, text="Search query:").grid(row=0, column=0, sticky="w")
        self.query = tk.StringVar(value="restaurants")
        ttk.Entry(top, textvariable=self.query, width=30).grid(row=0, column=1, padx=5)

        ttk.Label(top, text="Cities / countries (one per line):").grid(row=0, column=2, sticky="w")
        self.locations = tk.Text(top, width=28, height=3)
        self.locations.insert("1.0", "Shanghai, China")
        self.locations.grid(row=0, column=3, padx=5, rowspan=2)

        ttk.Label(top, text="Maximum results per location:").grid(row=1, column=0, sticky="w")
        self.limit = tk.StringVar(value="50")
        ttk.Entry(top, textvariable=self.limit, width=8).grid(row=1, column=1, sticky="w", padx=5)

        self.email_scan = tk.BooleanVar(value=False)
        ttk.Checkbutton(
            top, text="Find emails on websites (slower)",
            variable=self.email_scan
        ).grid(row=2, column=0, columnspan=2, sticky="w", pady=5)

        buttons = ttk.Frame(top)
        buttons.grid(row=2, column=2, columnspan=2, sticky="e")
        self.start_btn = ttk.Button(buttons, text="Start scraping", command=self.start)
        self.start_btn.pack(side="left", padx=3)
        self.stop_btn = ttk.Button(buttons, text="Stop", command=self.stop, state="disabled")
        self.stop_btn.pack(side="left", padx=3)
        ttk.Button(buttons, text="Refresh table", command=self.refresh_table).pack(side="left", padx=3)
        ttk.Button(buttons, text="Export", command=self.export).pack(side="left", padx=3)
        ttk.Button(buttons, text="Clear database", command=self.clear_database).pack(side="left", padx=3)

        fields_frame = ttk.LabelFrame(self.root, text="Output columns", padding=6)
        fields_frame.pack(fill="x", padx=8, pady=4)
        for i, (key, label) in enumerate(FIELDS.items()):
            var = tk.BooleanVar(value=True)
            self.selected_fields[key] = var
            ttk.Checkbutton(fields_frame, text=label, variable=var).grid(
                row=i // 7, column=i % 7, sticky="w", padx=5
            )

        self.status = tk.StringVar(value="Ready")
        ttk.Label(self.root, textvariable=self.status).pack(fill="x", padx=10, pady=3)

        columns = list(FIELDS.keys())
        self.tree = ttk.Treeview(self.root, columns=columns, show="headings")
        for key, label in FIELDS.items():
            self.tree.heading(key, text=label)
            self.tree.column(key, width=150, minwidth=80)
        yscroll = ttk.Scrollbar(self.root, orient="vertical", command=self.tree.yview)
        xscroll = ttk.Scrollbar(self.root, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        self.tree.pack(fill="both", expand=True, side="top", padx=(8, 0), pady=(4, 0))
        xscroll.pack(fill="x", padx=8)
        yscroll.place(relx=0.985, rely=0.42, relheight=0.48, anchor="ne")

        self.log = tk.Text(self.root, height=5, wrap="word")
        self.log.pack(fill="x", padx=8, pady=6)

    def log_msg(self, text):
        self.log.insert("end", text + "\n")
        self.log.see("end")

    def start(self):
        if self.running:
            return
        query = clean(self.query.get())
        locations = [
            clean(x) for x in self.locations.get("1.0", "end").splitlines()
            if clean(x)
        ]
        try:
            limit = int(self.limit.get())
            if limit < 1:
                raise ValueError
        except ValueError:
            messagebox.showerror("Error", "Maximum results must be a number greater than zero.")
            return
        if not query or not locations:
            messagebox.showerror("Error", "Enter a search query and at least one location.")
            return

        self.running = True
        self.stop_event.clear()
        self.start_btn.config(state="disabled")
        self.stop_btn.config(state="normal")
        threading.Thread(
            target=self.worker,
            args=(query, locations, limit, self.email_scan.get()),
            daemon=True
        ).start()

    def stop(self):
        self.stop_event.set()
        self.status.set("Stopping…")

    def worker(self, query, locations, limit, scan_email):
        try:
            # Import here so the app can still open and show a useful error
            # if Playwright has not been installed yet.
            from playwright.sync_api import sync_playwright

            with sync_playwright() as p:
                # Opens a normal browser window so you can complete any consent prompts.
                browser = p.chromium.launch(headless=False)
                context = browser.new_context(locale="en-US")
                page = context.new_page()

                for location in locations:
                    if self.stop_event.is_set():
                        break

                    search_term = f"{query} in {location}"
                    search = "https://www.google.com/maps/search/" + quote_plus(search_term)
                    self.events.put(("log", f"Searching: {query} — {location}"))
                    page.goto(search, wait_until="domcontentloaded", timeout=60000)
                    page.wait_for_timeout(2500)

                    # Collect result links progressively while scrolling the results feed.
                    urls = []
                    seen = set()
                    feed = page.locator('div[role="feed"]').first
                    if feed.count() == 0:
                        self.events.put(("log", "Results feed not found; a consent page or CAPTCHA may be blocking access."))
                        continue

                    no_new_rounds = 0
                    while len(urls) < limit and no_new_rounds < 5 and not self.stop_event.is_set():
                        links = page.locator('a[href*="/maps/place/"]')
                        count_before = len(urls)
                        for i in range(links.count()):
                            try:
                                href = links.nth(i).get_attribute("href", timeout=1000)
                                label = clean(links.nth(i).get_attribute("aria-label", timeout=1000) or "")
                                if href and "/maps/place/" in href:
                                    full = page.url if False else href
                                    full = full if full.startswith("http") else "https://www.google.com" + full
                                    key = normalize_url(full)
                                    if key not in seen:
                                        seen.add(key)
                                        urls.append((full, label))
                                        self.events.put(("status", f"{location}: found {len(urls)} / {limit} links"))
                                        if len(urls) >= limit:
                                            break
                            except Exception:
                                continue

                        if len(urls) == count_before:
                            no_new_rounds += 1
                        else:
                            no_new_rounds = 0

                        try:
                            feed.evaluate("(el) => el.scrollBy(0, Math.max(500, el.clientHeight * 0.8))")
                        except Exception:
                            pass
                        page.wait_for_timeout(900)

                    self.events.put(("log", f"Found {len(urls)} results to process."))

                    for index, (url, label) in enumerate(urls[:limit], start=1):
                        if self.stop_event.is_set():
                            break
                        try:
                            page.goto(url, wait_until="domcontentloaded", timeout=45000)
                            try:
                                page.locator("h1").first.wait_for(timeout=7000)
                            except Exception:
                                pass
                            page.wait_for_timeout(400)
                            record = extract_details(page, label)

                            if scan_email and record.get("website"):
                                try:
                                    website = record["website"]
                                    if website.startswith("http"):
                                        site_page = context.new_page()
                                        site_page.goto(website, wait_until="domcontentloaded", timeout=12000)
                                        site_page.wait_for_timeout(300)
                                        record["email"] = record.get("email") or extract_emails(
                                            site_page.locator("body").inner_text(timeout=3000)
                                        )
                                        # Also check mailto links.
                                        mailtos = site_page.locator('a[href^="mailto:"]')
                                        found = []
                                        for j in range(min(mailtos.count(), 20)):
                                            href = mailtos.nth(j).get_attribute("href") or ""
                                            found.extend(re.findall(
                                                r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
                                                href
                                            ))
                                        if found:
                                            record["email"] = ", ".join(sorted(set(
                                                (record.get("email", "").split(", ") if record.get("email") else []) + found
                                            )))
                                        site_page.close()
                                except Exception:
                                    pass

                            result = save_record(record)
                            self.events.put(("record", record, result))
                            self.events.put((
                                "status",
                                f"{location} | {index}/{min(len(urls), limit)} | {record.get('name') or label} | {result}"
                            ))
                        except Exception as exc:
                            self.events.put(("log", f"Error on result {index}: {exc}"))
                            continue

                try:
                    browser.close()
                except Exception:
                    pass

        except Exception as exc:
            self.events.put(("log", f"Main error: {exc}"))
        finally:
            self.events.put(("finished",))

    def process_events(self):
        try:
            while True:
                item = self.events.get_nowait()
                kind = item[0]
                if kind == "log":
                    self.log_msg(item[1])
                elif kind == "status":
                    self.status.set(item[1])
                elif kind == "record":
                    record, result = item[1], item[2]
                    if result in ("new", "updated"):
                        self.add_or_update_row(record)
                    self.log_msg(f"{result}: {record.get('name') or 'Unnamed'}")
                elif kind == "finished":
                    self.running = False
                    self.start_btn.config(state="normal")
                    self.stop_btn.config(state="disabled")
                    self.status.set("Finished")
                    self.refresh_table()
        except queue.Empty:
            pass
        self.root.after(150, self.process_events)

    def add_or_update_row(self, record):
        url = normalize_url(record.get("maps_url", ""))
        values = [record.get(k, "") for k in FIELDS]
        for iid in self.tree.get_children():
            vals = self.tree.item(iid, "values")
            if len(vals) == len(FIELDS) and normalize_url(vals[list(FIELDS).index("maps_url")]) == url:
                self.tree.item(iid, values=values)
                return
        self.tree.insert("", "end", values=values)

    def refresh_table(self):
        for iid in self.tree.get_children():
            self.tree.delete(iid)
        with sqlite3.connect(DB_FILE) as con:
            con.row_factory = sqlite3.Row
            rows = con.execute("SELECT * FROM leads ORDER BY id DESC").fetchall()
        for row in rows:
            self.tree.insert("", "end", values=[row[k] or "" for k in FIELDS])

    def clear_database(self):
        if self.running:
            messagebox.showwarning("Cannot clear database", "Stop scraping and wait for it to finish first.")
            return
        confirm = messagebox.askyesno(
            "Confirm data deletion",
            "All records in the database will be permanently deleted. Continue?"
        )
        if not confirm:
            return
        try:
            with sqlite3.connect(DB_FILE) as con:
                con.execute("DELETE FROM leads")
                # شناسه‌ها هم برای مجموعه‌ی جدید از ابتدا شروع شوند.
                con.execute("DELETE FROM sqlite_sequence WHERE name='leads'")
            self.refresh_table()
            self.status.set("Database cleared; ready to collect a new list.")
            self.log_msg("All previous records have been deleted from the database.")
            messagebox.showinfo("Done", "Previous data cleared. You can now collect a new list.")
        except Exception as e:
            messagebox.showerror("Error", f"Could not clear database:\n{e}")

    def export(self):
        keys = [k for k, var in self.selected_fields.items() if var.get()]
        if not keys:
            messagebox.showwarning("Export", "Select at least one column.")
            return

        path = filedialog.asksaveasfilename(
            title="Save export",
            defaultextension=".csv",
            filetypes=[
                ("CSV", "*.csv"),
                ("JSON", "*.json"),
                ("Excel", "*.xlsx"),
                ("Text", "*.txt"),
            ]
        )
        if not path:
            return

        with sqlite3.connect(DB_FILE) as con:
            con.row_factory = sqlite3.Row
            rows = [dict(r) for r in con.execute("SELECT * FROM leads ORDER BY id DESC")]

        data = [{k: row.get(k, "") for k in keys} for row in rows]
        try:
            suffix = Path(path).suffix.lower()
            if suffix == ".json":
                Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
            elif suffix == ".txt":
                with open(path, "w", encoding="utf-8") as f:
                    for row in data:
                        f.write(" | ".join(f"{FIELDS[k]}: {row[k]}" for k in keys) + "\n")
            elif suffix == ".xlsx":
                try:
                    from openpyxl import Workbook
                except ImportError:
                    messagebox.showerror("Installation required", "To export to Excel, run:\npip install openpyxl")
                    return
                wb = Workbook()
                ws = wb.active
                ws.title = "Leads"
                ws.append([FIELDS[k] for k in keys])
                for row in data:
                    ws.append([row[k] for k in keys])
                wb.save(path)
            else:
                with open(path, "w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.writer(f)
                    writer.writerow([FIELDS[k] for k in keys])
                    for row in data:
                        writer.writerow([row[k] for k in keys])
            messagebox.showinfo("Export", f"Saved to:\n{path}")
        except Exception as exc:
            messagebox.showerror("Export error", str(exc))


if __name__ == "__main__":
    root = tk.Tk()
    app = MapsApp(root)
    root.mainloop()